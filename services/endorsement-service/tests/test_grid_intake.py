import io

import openpyxl
import pytest
from django.core.files.uploadedfile import SimpleUploadedFile

from apps.computation.policy_client import PolicyUnavailable
from apps.intake.models import GuiRow, GuiSession
from apps.intake.schema import is_auto_confirmable, suggest_mapping
from apps.jobs.models import EndorsementJob

from .conftest import BROKER_A, BROKER_B, EMP_A, client_for

pytestmark = pytest.mark.django_db(transaction=True)
B = "/api/v1/broker/endorsements/"


def plan(c=None, prefix=B, **kw):
    body = {"policy_id": 100, "op_type": "ADD", "employee_count": 3, "effective_date": "2030-07-01", **kw}
    return (c or client_for()).post(prefix + "intake/plan/", body, format="json")


def grid(c, sid, rows, op="ADD", **kw):
    return c.post(f"{B}gui/sessions/{sid}/grid/validate/", {"op_type": op, "rows": rows, **kw}, format="json")


def row(code, rel, first, dob, **kw):
    return {"employee_code": code, "relation": rel, "first_name": first, "dob": dob, "gender": "F", "sum_insured": "500000", **kw}


# ── how many employees decides how they are entered ───────────────────────────────────────────────────────────────────────
@pytest.mark.parametrize("count,mode", [(1, "GRID"), (99, "GRID"), (100, "GRID"), (101, "FILE"), (5000, "FILE")])
def test_the_count_picks_grid_or_file(world, count, mode):
    r = plan(employee_count=count)
    body = r.json()
    assert body["mode"] == mode and body["employee_count"] == count and body["grid_max_employees"] == 100
    if mode == "GRID":
        assert r.status_code == 201 and body["session_id"] and body["file"] is None
        assert body["grid"]["initial_rows"] == count and body["grid"]["row_cap"] == count * 8
        assert {c["field"] for c in body["schema"]["columns"]} >= {"employee_code", "relation", "dob"}
        s = GuiSession.objects.get(pk=body["session_id"])
        assert s.expected_employees == count and str(s.effective_date) == "2030-07-01"
    else:
        assert r.status_code == 200 and body["session_id"] is None and body["grid"] is None
        assert body["file"]["estimated_chunks"] == -(-count // 500) and "xlsx" in body["file"]["formats"]
        assert GuiSession.objects.count() == 0                      # a file needs no session


def test_the_threshold_is_a_setting(world, settings):
    settings.GRID_MAX_EMPLOYEES = 10
    assert plan(employee_count=10).json()["mode"] == "GRID" and plan(employee_count=11).json()["mode"] == "FILE"


@pytest.mark.parametrize("kw,code", [({"employee_count": 0}, 400), ({"employee_count": "many"}, 400), ({"employee_count": None}, 400), ({"op_type": "MERGE"}, 400),
                                     ({"effective_date": "07/01/2030"}, 400), ({"policy_id": 102}, 422), ({"policy_id": 999}, 404), ({"policy_id": 200}, 404)])
def test_plan_refuses_bad_input_and_foreign_policies(world, kw, code):
    assert plan(**kw).status_code == code


def test_who_may_plan_is_who_may_submit(world):
    assert plan(client_for("BROKER_NON_ADMIN")).status_code == 403
    assert plan(client_for("EMPLOYER_NON_ADMIN", broker_id=BROKER_A, employer_id=EMP_A), prefix="/api/v1/employer/endorsements/").status_code == 201
    assert plan(client_for("BROKER_ADMIN", broker_id=BROKER_B)).status_code == 404


# ── the live grid ──────────────────────────────────────────────────────────────────────────────────────────────────────────
def test_grid_check_sees_across_rows_and_stores_nothing(world):
    c = client_for()
    sid = plan(c).json()["session_id"]
    rows = [row("1001", "SELF", "Asha", "17/05/1985", pan="ABCDE1234F"),
            row("1001", "SPOUSE", "Ravi", "02/02/1983"),
            {},                                                         # an empty row on screen is ignored, not an error
            row("1001", "SPOUSE", "Ravi again", "02/02/1983"),          # the same person twice
            row("1002", "SELF", "Meera", "03/03/1990", pan="ABCDE1234F"),   # a PAN already used in this grid
            row("1003", "SELF", "Raj", "31/02/1980")]                   # not a date
    body = grid(c, sid, rows).json()
    rules = {i: {x["rule"] for x in r["issues"]} for i, r in enumerate(body["rows"])}
    assert body["rows"][2]["blank"] and not body["rows"][2]["issues"] and not body["rows"][2]["valid"]
    assert "DUPLICATE_ROW_IN_FILE" in rules[3] and "DUPLICATE_PAN_IN_FILE" in rules[4] and "INVALID_DOB" in rules[5]
    assert body["rows"][0]["valid"] and body["rows"][1]["valid"]
    assert GuiRow.objects.count() == 0                                  # checking is not saving


def test_grid_check_runs_the_family_rules_in_one_call(world):
    c = client_for()
    sid = plan(c).json()["session_id"]
    kids = [row("1001", "SELF", "Asha", "17/05/1985")] + [row("1001", rel, f"Kid{i}", f"0{i}/01/201{i}") for i, rel in enumerate(["SON", "SON_II", "DAUGHTER"], start=1)]
    world.calls.clear()
    body = grid(c, sid, kids).json()
    assert world.calls.count("validate_families") == 1                  # the whole grid, one round trip
    assert body["checks"]["family_rules"] is True
    assert any(i["severity"] == "ERROR" for r in body["rows"] for i in r["issues"])           # policy 100 allows two children


def test_family_service_down_does_not_blank_the_grid(world):
    c = client_for()
    sid = plan(c).json()["session_id"]
    world.fail_next.append(PolicyUnavailable("down"))
    body = grid(c, sid, [row("1001", "SELF", "Asha", "17/05/1985"), row("1001", "SPOUSE", "Ravi", "02/02/1983"), row("1002", "SELF", "Meera", "bad")]).json()
    assert body["checks"]["family_rules"] is False                       # said so, instead of pretending
    assert "INVALID_DOB" in {i["rule"] for i in body["rows"][2]["issues"]} and body["rows"][0]["valid"]


def test_the_summary_is_pandas_over_the_grid(world):
    c = client_for()
    sid = plan(c).json()["session_id"]
    rows = [row("1001", "SELF", "Asha", "17/05/1985"), row("1001", "SPOUSE", "Ravi", "02/02/1983", sum_insured="500000"),
            row("1001", "DAUGHTER", "Tina", "01/01/2015", sum_insured="300000"), row("1002", "SELF", "Meera", "03/03/1960", sum_insured="200000"), row("1003", "SELF", "Raj", "x")]
    s = grid(c, sid, rows).json()["summary"]
    assert (s["rows_typed"], s["employees_entered"], s["expected_employees"]) == (5, 3, 3)
    assert s["rows_with_errors"] == 1 and s["by_rule"][0]["rule"] == "INVALID_DOB"
    assert s["relations"] == {"SELF": 3, "SPOUSE": 1, "DAUGHTER": 1}
    assert s["family_sizes"] == {"1": 2, "3": 1}
    assert s["ages"]["max"] == 70 and s["ages"]["min"] == 15            # on the effective date, 1 Jul 2030
    assert s["ages_by_band"]["66+"] == 1
    assert s["sum_insured"]["count"] == 5 and s["sum_insured"]["max"] == 500000 and s["sum_insured"]["total"] == 2000000.0
    assert s["completeness"]["dob"] == 80.0


def test_the_grid_has_a_row_cap_from_the_count(world):
    c = client_for()
    sid = plan(c, employee_count=2).json()["session_id"]
    assert grid(c, sid, [row("1", "SELF", "A", "01/01/1990")] * 16).status_code == 200
    assert grid(c, sid, [row("1", "SELF", "A", "01/01/1990")] * 17).status_code == 422


def test_grid_check_is_scoped_to_the_callers_session(world):
    sid = plan().json()["session_id"]
    assert grid(client_for(broker_id=BROKER_B), sid, [{}]).status_code == 404
    assert grid(client_for(), sid, "nope").status_code == 400


# ── save and submit ────────────────────────────────────────────────────────────────────────────────────────────────────────
def test_save_then_submit_makes_a_job(world):
    c = client_for()
    sid = plan(c).json()["session_id"]
    rows = [row("1001", "SELF", "Asha", "17/05/1985"), {}, row("1002", "SELF", "Meera", "03/03/1990")]
    r = c.put(f"{B}gui/sessions/{sid}/grid/", {"op_type": "ADD", "rows": rows}, format="json")
    assert r.status_code == 200 and r.json()["saved_rows"] == 2
    assert GuiRow.objects.filter(session_id=sid).count() == 2
    again = c.put(f"{B}gui/sessions/{sid}/grid/", {"op_type": "ADD", "rows": rows[:1]}, format="json")       # saving replaces, never piles up
    assert again.json()["saved_rows"] == 1 and GuiRow.objects.filter(session_id=sid).count() == 1
    sub = c.post(f"{B}gui/sessions/{sid}/submit/")
    assert sub.status_code == 202 and sub.json()["jobs"][0]["rows"] == 1
    assert EndorsementJob.objects.get(pk=sub.json()["jobs"][0]["job_id"]).source_type == "GUI"


def test_a_grid_with_errors_saves_but_will_not_submit(world):
    c = client_for()
    sid = plan(c).json()["session_id"]
    c.put(f"{B}gui/sessions/{sid}/grid/", {"op_type": "ADD", "rows": [row("1001", "SELF", "Asha", "nope")]}, format="json")
    assert GuiRow.objects.filter(session_id=sid).count() == 1           # nothing typed is lost
    assert c.post(f"{B}gui/sessions/{sid}/submit/").status_code == 422


# ── the grid and the file path are one rule book ───────────────────────────────────────────────────────────────────────────
def test_grid_and_csv_agree_on_which_rows_fail(world):
    c = client_for()
    lines = [("1001", "Self", "Asha Rao", "17/05/1985", "F", "500000", "ABCDE1234F"), ("1001", "Spouse", "Ravi Rao", "02-Feb-1983", "M", "500000", ""),
             ("1002", "Self", "Meera", "not-a-date", "F", "300000", ""), ("1003", "Cousin", "Odd", "01/01/2000", "M", "100000", ""),
             ("1004", "Self", "Raj", "07/07/1980", "M", "-5", ""), ("1001", "Spouse", "Dup", "02-Feb-1983", "M", "500000", ""),
             ("1005", "Self", "Pan1", "01/01/1990", "M", "100", "PQRST5678K"), ("1006", "Self", "Pan2", "01/01/1991", "M", "100", "PQRST5678K")]
    sid = plan(c, employee_count=8).json()["session_id"]
    cells = [{"employee_code": a, "relation": b, "full_name": n, "dob": d, "gender": g, "sum_insured": si, "pan": pan} for a, b, n, d, g, si, pan in lines]
    in_grid = {i for i, r in enumerate(grid(c, sid, cells).json()["rows"]) if not r["valid"]}
    text = "Emp Code,Relationship,Member Name,DOB,Gender,Sum Insured,PAN,Effective Date\n" + "\n".join(",".join(x) + ",2030-07-01" for x in lines) + "\n"
    up = c.post(B + "csv/upload/", {"file": SimpleUploadedFile("f.csv", text.encode(), content_type="text/csv"), "op_type": "ADD", "policy_id": 100, "effective_date": "2030-07-01"}, format="multipart")
    errs = c.get(f"{B}{up.json()['job_id']}/errors/?page_size=100").json()
    in_file = {e["line"] - 2 for e in errs["results"] if e["severity"] == "ERROR"}
    assert in_grid == in_file and in_grid                              # same rows, same verdict


# ── the spreadsheet for more than 100 ──────────────────────────────────────────────────────────────────────────────────────
def test_template_downloads_and_our_reader_understands_every_column(world):
    r = client_for().get(B + "intake/template/", {"policy_id": 100, "op_type": "ADD", "employee_count": 250, "effective_date": "2030-07-01"})
    assert r.status_code == 200 and r["Content-Type"].endswith("spreadsheetml.sheet") and "attachment" in r["Content-Disposition"]
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    ws = wb["Endorsement"]
    assert ws.max_row == 251 and [c.value for c in ws[2]][1] == "SELF" and ws["N2"].value == "2030-07-01"
    sug = suggest_mapping([c.value for c in ws[1]])
    assert all(s.confidence == 1.0 and s.field for s in sug) and is_auto_confirmable(sug, "ADD")
    assert len(ws.data_validations.dataValidation) >= 5 and "Instructions" in wb.sheetnames


def test_a_filled_template_goes_straight_through_the_file_path(world):
    c = client_for()
    wb = openpyxl.load_workbook(io.BytesIO(c.get(B + "intake/template/", {"policy_id": 100, "op_type": "ADD", "employee_count": 2, "effective_date": "2030-07-01"}).content))
    ws = wb["Endorsement"]
    for r, (code, first, dob) in enumerate([("7001", "Asha", "1985-05-17"), ("7002", "Meera", "1990-03-03")], start=2):
        ws.cell(r, 1, code); ws.cell(r, 3, first); ws.cell(r, 5, dob); ws.cell(r, 13, 500000)
    buf = io.BytesIO(); wb.save(buf)
    up = c.post(B + "csv/upload/", {"file": SimpleUploadedFile("t.xlsx", buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), "op_type": "ADD", "policy_id": 100}, format="multipart")
    assert up.status_code == 202, up.content
    job = EndorsementJob.objects.get(pk=up.json()["job_id"])
    assert job.status == "VALIDATION_PASSED" and job.total_rows == 2 and job.failed_rows == 0


@pytest.mark.parametrize("q,code", [({"op_type": "ADD", "employee_count": 5}, 400), ({"policy_id": 100, "op_type": "X", "employee_count": 5}, 400), ({"policy_id": 102, "op_type": "ADD", "employee_count": 5}, 422)])
def test_template_input_is_checked(world, q, code):
    assert client_for().get(B + "intake/template/", q).status_code == code
