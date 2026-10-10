import io

import openpyxl
import pytest

from apps.jobs.models import EndorsementJob, EndorsementMember

from .conftest import BROKER_B, client_for, csv_text, upload

pytestmark = pytest.mark.django_db(transaction=True)
BASE = "/api/v1/broker/endorsements/"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIXED = csv_text("1,Self,Good,1990-01-01,M,500000,,", "2,Self,,1990-01-01,M,500000,,", "3,Self,Also Good,1990-01-01,M,500000,,")


def job_from(resp):
    assert resp.status_code in (200, 202), resp.content
    return EndorsementJob.objects.get(pk=resp.json()["job_id"])


def sheet(resp):
    return openpyxl.load_workbook(io.BytesIO(resp.content))


def test_a_partly_failed_upload_carries_on_with_the_valid_rows_by_itself(flow):
    job = job_from(upload(client_for(), MIXED, effective_date="2030-04-01"))
    assert job.status == "ENDORSED"
    assert sorted(EndorsementMember.objects.values_list("employee_code", flat=True)) == ["1", "3"]
    assert job.failed_rows == 1 and job.total_rows == 3                                         # the failure is still on the record
    errs = client_for().get(f"{BASE}{job.id}/errors/").json()
    assert errs["count"] == 1 and errs["results"][0]["employee_code"] == "2"


def test_it_can_still_be_held_for_a_person_when_the_setting_is_off(flow, settings):
    settings.PARTIAL_AUTO_CONTINUE = False
    job = job_from(upload(client_for(), MIXED, effective_date="2030-04-01"))
    assert job.status == "VALIDATION_PARTIAL" and EndorsementMember.objects.count() == 0


def test_a_fully_failed_upload_stops_and_does_not_continue(flow):
    job = job_from(upload(client_for(), csv_text("1,Self,,1990-01-01,M,500000,,"), effective_date="2030-04-01"))
    assert job.status == "VALIDATION_FAILED" and EndorsementMember.objects.count() == 0


def test_the_result_sheet_has_the_legacy_layout_and_one_row_per_uploaded_row(flow):
    from apps.validation import legacy_layouts
    job = job_from(upload(client_for(), MIXED, effective_date="2030-04-01"))
    r = client_for().get(f"{BASE}{job.id}/errors/download/")
    assert r.status_code == 200 and r["Content-Type"] == XLSX
    assert r["Content-Disposition"] == 'attachment; filename="AddMemberError.xlsx"'
    wb = sheet(r)
    assert wb.sheetnames == ["Worksheet"]
    rows = list(wb["Worksheet"].iter_rows(values_only=True))
    header = list(rows[0])
    assert header == legacy_layouts.ADD + ["Status", "Reason"] and len(header) == 108
    assert len(rows) == 4                                                                       # every uploaded row, good or bad
    by_code = {dict(zip(header, r))["Employee Code"]: dict(zip(header, r)) for r in rows[1:]}
    assert by_code["1"]["Status"] == "Member added successfully." and by_code["1"]["Reason"] is None
    assert by_code["3"]["Status"] == "Member added successfully."
    assert by_code["2"]["Status"] == "Failed to add." and by_code["2"]["Reason"]
    assert by_code["1"]["Employee First Name"] == "Good" and by_code["1"]["Employee DOB"] == "01-01-1990" and by_code["1"]["Suminsured"] == 500000
    assert by_code["1"]["Relationship with Employee"] == "Self" and by_code["1"]["Inception/Endorsement Date"] == "01-04-2030"


def test_a_row_that_passed_but_was_not_endorsed_is_not_reported_as_added(flow, settings):
    settings.PARTIAL_AUTO_CONTINUE = False
    job = job_from(upload(client_for(), MIXED, effective_date="2030-04-01"))
    rows = list(sheet(client_for().get(f"{BASE}{job.id}/errors/download/"))["Worksheet"].iter_rows(values_only=True))
    h = list(rows[0])
    status = {dict(zip(h, r))["Employee Code"]: dict(zip(h, r))["Status"] for r in rows[1:]}
    assert status == {"1": "Not yet added.", "2": "Failed to add.", "3": "Not yet added."}


def test_a_clean_job_reports_every_row_as_added(flow):
    job = job_from(upload(client_for(), csv_text("1,Self,A,1990-01-01,M,500000,,"), effective_date="2030-04-01"))
    rows = list(sheet(client_for().get(f"{BASE}{job.id}/errors/download/"))["Worksheet"].iter_rows(values_only=True))
    assert len(rows) == 2 and rows[1][-2] == "Member added successfully."


def test_values_that_look_like_formulas_are_written_as_text(flow):
    text = csv_text("1,Self,=HYPERLINK(\"http://x\"),1990-01-01,M,500000,,", "2,Self,Fine,1990-01-01,M,500000,,")
    job = job_from(upload(client_for(), text, effective_date="2030-04-01"))
    wb = sheet(client_for().get(f"{BASE}{job.id}/errors/download/"))
    for sh in wb:
        for row in sh.iter_rows(values_only=True):
            assert not any(isinstance(v, str) and v.startswith("=") for v in row), row


def test_another_broker_cannot_download_it_and_a_stranger_gets_nothing(flow):
    job = job_from(upload(client_for(), MIXED, effective_date="2030-04-01"))
    assert client_for("BROKER_ADMIN", BROKER_B).get(f"{BASE}{job.id}/errors/download/").status_code == 404
    from rest_framework.test import APIClient
    assert APIClient().get(f"{BASE}{job.id}/errors/download/").status_code in (401, 403)
