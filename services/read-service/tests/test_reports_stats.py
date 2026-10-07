import csv
import io
from datetime import UTC, datetime, timedelta

import pytest

from apps.core import mongo
from apps.projections.handlers import handle_event

from .conftest import BROKER_A, BROKER_B, EMP_A, client_for, sample

BASE = "/api/v1/broker/read/"


def job(jid, state="ENDORSED", ts="2030-10-01T10:00:00+00:00", op="ADD", broker=BROKER_A, employer=EMP_A, policy=1, terminal=True):
    e = {**sample("endorsement-state-changed"), "event_id": f"ev-{jid}-{state}-{ts}", "job_id": jid, "to_state": state, "timestamp": ts, "op_type": op,
         "broker_id": broker, "employer_id": employer, "policy_id": policy, "is_terminal": terminal}
    return e


def applied(jid, members, op="ADD", broker=BROKER_A, employer=EMP_A, policy=1):
    return {**sample("endorsement-applied"), "event_id": f"ap-{jid}", "job_id": jid, "members_changed": members, "op_type": op, "broker_id": broker,
            "employer_id": employer, "policy_id": policy}


@pytest.fixture
def world(db):
    for jid, ts, op, n in (("j1", "2030-09-10T10:00:00+00:00", "ADD", 3), ("j2", "2030-09-20T10:00:00+00:00", "REMOVE", 1),
                           ("j3", "2030-10-02T10:00:00+00:00", "ADD", 5), ("j4", "2030-10-03T10:00:00+00:00", "UPDATE", 2)):
        handle_event(db, job(jid, ts=ts, op=op))
        handle_event(db, applied(jid, n, op))
    handle_event(db, job("other", broker=BROKER_B, employer=20, policy=9, ts="2030-10-05T10:00:00+00:00"))
    handle_event(db, applied("other", 99, broker=BROKER_B, employer=20, policy=9))
    return db


def test_monthly_activity_buckets_members_by_operation_and_month(world):
    r = client_for("BROKER_ADMIN", BROKER_A).get(BASE + "dashboard/endorsement-activity/").json()
    assert r["months"] == [{"month": "2030-09", "added": 3, "updated": 0, "removed": 1, "endorsements": 2},
                           {"month": "2030-10", "added": 5, "updated": 2, "removed": 0, "endorsements": 2}]          # the other broker's 99 is not here
    r = client_for("BROKER_ADMIN", BROKER_A).get(BASE + "dashboard/endorsement-activity/?from=2030-10-01&to=2030-10-02").json()
    assert [m["month"] for m in r["months"]] == ["2030-10"] and r["months"][0]["added"] == 5 and r["months"][0]["updated"] == 0
    c = client_for("BROKER_ADMIN", BROKER_A)
    assert c.get(BASE + "dashboard/endorsement-activity/?from=yesterday").status_code == 400
    assert c.get(BASE + "dashboard/endorsement-activity/?policy_id=x").status_code == 400


def test_endorsed_at_keeps_the_earliest_endorsed_transition_even_when_events_arrive_late(db):
    handle_event(db, job("j", ts="2030-10-05T10:00:00+00:00"))
    handle_event(db, job("j", ts="2030-10-02T10:00:00+00:00"))
    assert db[mongo.JOBS].find_one({"_id": "j"})["endorsed_at"].day == 2


def test_red_flags_cover_low_deposits_stuck_and_blocked_jobs_and_failed_syncs(db):
    now = datetime.now(UTC)
    old, recent = (now - timedelta(days=20)).isoformat(), (now - timedelta(hours=2)).isoformat()
    handle_event(db, job("stuck", state="VALIDATION_IN_PROGRESS", ts=old, terminal=False))
    handle_event(db, job("blocked", state="CD_INSUFFICIENT", ts=recent, terminal=False))
    handle_event(db, job("fine", state="PREMIUM_CALCULATED", ts=recent, terminal=False))
    handle_event(db, job("done", state="ENDORSED", ts=old, terminal=True))
    handle_event(db, sample("cd-balance-updated"))
    handle_event(db, sample("tpa-sync-failed"))
    handle_event(db, job("theirs", state="CD_INSUFFICIENT", ts=recent, terminal=False, broker=BROKER_B, employer=20, policy=9))
    r = client_for("BROKER_ADMIN", BROKER_A).get(BASE + "dashboard/red-flags/").json()
    kinds = {(f["type"], f.get("job_id") or f.get("sync_id") or f.get("policy_id"), f["priority"]) for f in r["flags"]}
    assert ("ENDORSEMENT", "stuck", "High") in kinds and ("ENDORSEMENT", "blocked", "Medium") in kinds and ("CD_BALANCE", 1, "Medium") in kinds
    assert not any(f.get("job_id") in ("fine", "done", "theirs") for f in r["flags"]) and any(f["type"] == "TPA_SYNC" for f in r["flags"])
    assert r["flags"][0]["priority"] == "High"                                                                    # sorted by priority
    assert client_for("BROKER_ADMIN", BROKER_A).get(BASE + "dashboard/red-flags/?pending_days=40").json()["count"] < r["count"]


def parse(resp):
    return list(csv.reader(io.StringIO(resp.content.decode("utf-8-sig"))))


def test_csv_report_is_scoped_filtered_and_excel_safe(world):
    r = client_for("BROKER_ADMIN", BROKER_A).get(BASE + "reports/endorsements/")
    assert r.status_code == 200 and r["Content-Type"].startswith("text/csv") and r["X-Report-Truncated"] == "false" and "no-store" in r["Cache-Control"]
    rows = parse(r)
    assert rows[0][:3] == ["job_id", "policy_id", "employer_id"] and sorted(x[0] for x in rows[1:]) == ["j1", "j2", "j3", "j4"]
    assert "other" not in r.content.decode()
    assert len(parse(client_for("BROKER_ADMIN", BROKER_A).get(BASE + "reports/endorsements/?op_type=ADD"))) == 3
    assert len(parse(client_for("BROKER_ADMIN", BROKER_A).get(BASE + "reports/endorsements/?from=2030-10-01&to=2030-10-02"))) == 2
    hr = client_for("EMPLOYER_ADMIN", BROKER_A, EMP_A).get("/api/v1/employer/read/reports/endorsements/")
    assert len(parse(hr)) == 5


def test_formula_injection_is_neutralised_in_csv_and_xlsx(db):
    db[mongo.POLICIES].insert_one({"_id": 5, "broker_id": BROKER_A, "employer_id": EMP_A, "policy_number": '=HYPERLINK("http://evil","x")', "status": "+cmd", "tpa_id": 3})
    c = client_for("BROKER_ADMIN", BROKER_A)
    row = parse(c.get(BASE + "reports/policies/"))[1]
    assert row[1].startswith("'=") and row[3].startswith("'+")
    import openpyxl
    x = c.get(BASE + "reports/policies/?file=xlsx")
    assert x["Content-Type"].endswith("spreadsheetml.sheet")
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    assert ws["B2"].value.startswith("'=") and ws["B2"].data_type == "s"


def test_report_validation_cap_and_permissions(world, settings):
    c = client_for("BROKER_ADMIN", BROKER_A)
    assert c.get(BASE + "reports/nope/").status_code == 404 and c.get(BASE + "reports/endorsements/?file=pdf").status_code == 400
    assert c.get(BASE + "reports/endorsements/?policy_id=x").status_code == 400 and c.get(BASE + "reports/endorsements/?from=31/12/2030").status_code == 400
    assert c.get(BASE + "reports/endorsements/?state[$ne]=x").status_code == 200                                    # unknown params are ignored, not queries
    settings.REPORT_MAX_ROWS = 2
    r = c.get(BASE + "reports/endorsements/")
    assert len(parse(r)) == 3 and r["X-Report-Truncated"] == "true"
    assert client_for("SYSTEM", None).get(BASE + "reports/endorsements/").status_code == 403
    assert client_for("EMPLOYER_ADMIN", BROKER_A, EMP_A).get(BASE + "reports/endorsements/").status_code == 403
